import os
import sys
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

import matplotlib
matplotlib.use('Agg')
matplotlib.rcParams['font.sans-serif'] = ['DejaVu Sans', 'Arial', 'Segoe UI']
matplotlib.rcParams['axes.unicode_minus'] = False
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(DOCS_DIR, 'images')
os.makedirs(IMG_DIR, exist_ok=True)

# -------------------------------------------------------------
# 1. HÀM TẠO CÁC HÌNH VẼ SƠ ĐỒ CHUYÊN NGHIỆP BẰNG MATPLOTLIB
# -------------------------------------------------------------

def draw_dfd_level_0():
    fig, ax = plt.subplots(figsize=(10, 5), dpi=300)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis('off')

    # Entity: Người dùng
    rect_user = patches.FancyBboxPatch((0.5, 1.8), 2.0, 1.4, boxstyle="round,pad=0.1", fc='#E8EAF6', ec='#3F51B5', lw=2)
    ax.add_patch(rect_user)
    ax.text(1.5, 2.5, "TÁC NHÂN\nNgười dùng / Sinh viên", ha='center', va='center', fontsize=11, fontweight='bold', color='#1A237E')

    # Process: Hệ thống StudyHub
    circle_sys = patches.Circle((5.0, 2.5), 1.3, fc='#E1F5FE', ec='#0288D1', lw=2.5)
    ax.add_patch(circle_sys)
    ax.text(5.0, 2.5, "HỆ THỐNG STUDYHUB\n(Kiến trúc Cashew\nLocal-First)", ha='center', va='center', fontsize=11, fontweight='bold', color='#01579B')

    # Data Store: SQLite & File Storage
    rect_store = patches.FancyBboxPatch((7.5, 1.8), 2.0, 1.4, boxstyle="round,pad=0.1", fc='#E8F5E9', ec='#388E3C', lw=2)
    ax.add_patch(rect_store)
    ax.text(8.5, 2.5, "KHO LƯU TRỮ\nDrift SQLite DB\n& File Storage", ha='center', va='center', fontsize=11, fontweight='bold', color='#1B5E20')

    # Arrows between User and System
    ax.annotate("Nhập liệu môn học,\nTải lên tệp tài liệu, Ghi chú", xy=(3.6, 2.9), xytext=(2.6, 2.9),
                arrowprops=dict(arrowstyle="->", color='#3F51B5', lw=1.8), fontsize=8.5, ha='center', color='#283593')
    ax.annotate("Hiển thị danh sách, Xem PDF,\nTải tệp về máy, Trạng thái", xy=(2.6, 2.1), xytext=(3.6, 2.1),
                arrowprops=dict(arrowstyle="->", color='#3F51B5', lw=1.8), fontsize=8.5, ha='center', color='#283593')

    # Arrows between System and Store
    ax.annotate("Ghi bản ghi, Tạo Tombstone,\nLưu trữ tệp cục bộ", xy=(7.4, 2.9), xytext=(6.4, 2.9),
                arrowprops=dict(arrowstyle="->", color='#0288D1', lw=1.8), fontsize=8.5, ha='center', color='#0277BD')
    ax.annotate("Truy vấn Reactive Stream,\nĐọc Base64 PDF, Metadata", xy=(6.4, 2.1), xytext=(7.4, 2.1),
                arrowprops=dict(arrowstyle="->", color='#0288D1', lw=1.8), fontsize=8.5, ha='center', color='#0277BD')

    plt.title("Sơ đồ luồng dữ liệu DFD Mức 0 (Context Diagram) - Ứng dụng StudyHub", fontsize=13, fontweight='bold', pad=15)
    out_path = os.path.join(IMG_DIR, "dfd_level_0.png")
    plt.tight_layout()
    plt.savefig(out_path, bbox_inches='tight')
    plt.close()
    return out_path

def draw_dfd_level_1():
    fig, ax = plt.subplots(figsize=(11, 7), dpi=300)
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 7)
    ax.axis('off')

    # Tác nhân Sinh viên
    rect_user = patches.FancyBboxPatch((0.5, 2.5), 1.8, 2.0, boxstyle="round,pad=0.1", fc='#E8EAF6', ec='#3F51B5', lw=2)
    ax.add_patch(rect_user)
    ax.text(1.4, 3.5, "TÁC NHÂN\nSinh viên\n(Người dùng)", ha='center', va='center', fontsize=10.5, fontweight='bold', color='#1A237E')

    # 4 Tiến trình chính (Processes)
    processes = [
        ("1.0 Quản lý\nMôn học", 4.2, 5.5, '#FFF3E0', '#FB8C00'),
        ("2.0 Quản lý &\nPhân loại Tài liệu", 4.2, 3.8, '#E1F5FE', '#039BE5'),
        ("3.0 Xem PDF Trực tiếp\n& Tải xuống", 4.2, 2.1, '#F3E5F5', '#8E24AA'),
        ("4.0 Tìm kiếm, Lọc\n& Đánh dấu sao", 4.2, 0.6, '#E8F5E9', '#43A047')
    ]

    for title, x, y, bg, border in processes:
        box = patches.FancyBboxPatch((x, y), 2.6, 1.0, boxstyle="round,pad=0.08", fc=bg, ec=border, lw=2)
        ax.add_patch(box)
        ax.text(x + 1.3, y + 0.5, title, ha='center', va='center', fontsize=9.5, fontweight='bold', color='#263238')

    # Kho dữ liệu (Data Stores)
    stores = [
        ("D1: Bảng Subjects\n(Danh mục Môn học)", 8.2, 5.5, '#FFF8E1', '#FFA000'),
        ("D2: Bảng Documents\n(Tài liệu & Metadata)", 8.2, 3.5, '#E0F7FA', '#00ACC1'),
        ("D3: Bảng DeleteLogs\n(Tombstone Xóa an toàn)", 8.2, 1.5, '#FFEBEE', '#E53935')
    ]

    for title, x, y, bg, border in stores:
        box = patches.Rectangle((x, y), 2.5, 0.9, fc=bg, ec=border, lw=2)
        ax.add_patch(box)
        ax.text(x + 1.25, y + 0.45, title, ha='center', va='center', fontsize=9, fontweight='bold', color='#37474F')

    # Đường mũi tên
    # Từ User tới 1.0, 2.0, 3.0, 4.0
    ax.annotate("", xy=(4.1, 6.0), xytext=(2.4, 4.2), arrowprops=dict(arrowstyle="->", color='#5C6BC0', lw=1.5))
    ax.annotate("", xy=(4.1, 4.3), xytext=(2.4, 3.7), arrowprops=dict(arrowstyle="->", color='#5C6BC0', lw=1.5))
    ax.annotate("", xy=(4.1, 2.6), xytext=(2.4, 3.2), arrowprops=dict(arrowstyle="->", color='#5C6BC0', lw=1.5))
    ax.annotate("", xy=(4.1, 1.1), xytext=(2.4, 2.8), arrowprops=dict(arrowstyle="->", color='#5C6BC0', lw=1.5))

    # Từ tiến trình tới Kho dữ liệu
    ax.annotate("", xy=(8.1, 6.0), xytext=(6.9, 6.0), arrowprops=dict(arrowstyle="<->", color='#FB8C00', lw=1.5))
    ax.annotate("", xy=(8.1, 4.0), xytext=(6.9, 4.3), arrowprops=dict(arrowstyle="<->", color='#039BE5', lw=1.5))
    ax.annotate("", xy=(8.1, 1.9), xytext=(6.9, 4.0), arrowprops=dict(arrowstyle="->", color='#E53935', lw=1.5))
    ax.annotate("", xy=(8.1, 3.8), xytext=(6.9, 2.6), arrowprops=dict(arrowstyle="<-", color='#8E24AA', lw=1.5))
    ax.annotate("", xy=(8.1, 3.6), xytext=(6.9, 1.1), arrowprops=dict(arrowstyle="<-", color='#43A047', lw=1.5))

    plt.title("Sơ đồ luồng dữ liệu DFD Mức 1 (Chi tiết 4 tiến trình) - StudyHub", fontsize=13, fontweight='bold', pad=15)
    out_path = os.path.join(IMG_DIR, "dfd_level_1.png")
    plt.tight_layout()
    plt.savefig(out_path, bbox_inches='tight')
    plt.close()
    return out_path

def draw_cashew_architecture():
    fig, ax = plt.subplots(figsize=(10, 7.5), dpi=300)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 8.5)
    ax.axis('off')

    layers = [
        ("TẦNG 1: PRESENTATION LAYER (Giao diện người dùng)\nMaterial 3 Responsive, Theme sáng/tối, GridView co giãn,\nDocumentsScreen, SubjectDetailScreen, FavoritesScreen, PdfViewerScreen", 7.0, '#E8EAF6', '#3F51B5'),
        ("TẦNG 2: STATE & CONTROLLER LAYER (Quản trị trạng thái)\nValueNotifier, Reactive Streams, AuthSyncService (Google Cloud Sync),\nBottomSheets & Dialogs, Tìm kiếm thời gian thực & Bộ lọc danh mục", 5.4, '#E0F2F1', '#00897B'),
        ("TẦNG 3: DOMAIN LOGIC & FORMATTERS (Quy tắc nghiệp vụ)\nNhận diện tự động đuôi file (PDF, Word, PPT), Phân loại Category,\nTính toán dung lượng byte, Đánh dấu sao quan trọng, Pastel Color Algorithm", 3.8, '#FFF3E0', '#FB8C00'),
        ("TẦNG 4: DATA ACCESS LAYER (DAOs & Storage Engine)\nSubjectDao, DocumentDao, FileStorageService (Base64 Web & Native Path),\nCơ chế Tombstone ghi nhận lịch sử xoá an toàn (DeleteLogs)", 2.2, '#E1F5FE', '#0288D1'),
        ("TẦNG 5: PERSISTENCE & LOCAL DATABASE (Lưu trữ cục bộ)\nDrift SQLite Engine (Local-First), Web IndexedDB / Android SQLite Native,\nHoạt động 100% Offline không phụ thuộc kết nối mạng", 0.6, '#F1F8E9', '#558B2F')
    ]

    for title, y, bg, border in layers:
        box = patches.FancyBboxPatch((0.5, y), 9.0, 1.25, boxstyle="round,pad=0.08", fc=bg, ec=border, lw=2.2)
        ax.add_patch(box)
        ax.text(5.0, y + 0.62, title, ha='center', va='center', fontsize=9.5, fontweight='bold', color='#212121')

    # Mũi tên 2 chiều giữa các tầng
    for y_arrow in [6.85, 5.25, 3.65, 2.05]:
        ax.annotate("", xy=(5.0, y_arrow - 0.15), xytext=(5.0, y_arrow + 0.15),
                    arrowprops=dict(arrowstyle="<->", color='#37474F', lw=2))

    plt.title("Sơ đồ Kiến trúc 5 Tầng Chuẩn Cashew (Cashew Architecture) - StudyHub", fontsize=13, fontweight='bold', pad=15)
    out_path = os.path.join(IMG_DIR, "cashew_architecture.png")
    plt.tight_layout()
    plt.savefig(out_path, bbox_inches='tight')
    plt.close()
    return out_path

def draw_erd():
    fig, ax = plt.subplots(figsize=(10, 6.5), dpi=300)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 7)
    ax.axis('off')

    # Entity 1: SUBJECTS
    sub_title = "BẢNG SUBJECTS (Môn học)"
    sub_fields = "\n".join([
        "PK: subjectPk (TEXT - UUID)",
        "name: TEXT (Tên môn học)",
        "code: TEXT (Mã môn, ví dụ CSE441)",
        "color: INTEGER (Mã màu HEX)",
        "icon: TEXT (Tên icon Material)",
        "dateCreated: DATETIME",
        "dateTimeModified: DATETIME"
    ])
    box1 = patches.FancyBboxPatch((0.5, 2.2), 3.8, 3.8, boxstyle="round,pad=0.08", fc='#EDE7F6', ec='#673AB7', lw=2)
    ax.add_patch(box1)
    ax.text(2.4, 5.6, sub_title, ha='center', va='center', fontsize=10.5, fontweight='bold', color='#4A148C')
    ax.plot([0.5, 4.3], [5.3, 5.3], color='#673AB7', lw=1.5)
    ax.text(0.7, 3.7, sub_fields, ha='left', va='center', fontsize=9, fontfamily='sans-serif', color='#212121')

    # Entity 2: DOCUMENTS
    doc_title = "BẢNG DOCUMENTS (Tài liệu)"
    doc_fields = "\n".join([
        "PK: documentPk (TEXT - UUID)",
        "name: TEXT (Tên tài liệu)",
        "FK: subjectFk (TEXT -> Subjects)",
        "type: INT (PDF, Word, PPT, Link)",
        "category: INT (Bài giảng, Tham khảo...)",
        "filePath: TEXT (Đường dẫn tệp máy)",
        "url: TEXT (Link web hoặc Base64)",
        "fileSize: INTEGER (Dung lượng bytes)",
        "note: TEXT (Ghi chú tài liệu)",
        "isPinned: BOOLEAN (Đánh dấu sao)",
        "dateCreated: DATETIME",
        "dateTimeModified: DATETIME"
    ])
    box2 = patches.FancyBboxPatch((5.7, 1.0), 3.8, 5.0, boxstyle="round,pad=0.08", fc='#E3F2FD', ec='#1976D2', lw=2)
    ax.add_patch(box2)
    ax.text(7.6, 5.6, doc_title, ha='center', va='center', fontsize=10.5, fontweight='bold', color='#0D47A1')
    ax.plot([5.7, 9.5], [5.3, 5.3], color='#1976D2', lw=1.5)
    ax.text(5.9, 3.1, doc_fields, ha='left', va='center', fontsize=8.5, fontfamily='sans-serif', color='#212121')

    # Entity 3: DELETELOGS
    del_title = "BẢNG DELETELOGS (Tombstone)"
    del_fields = "\n".join([
        "PK: logPk (TEXT - UUID)",
        "entryPk: TEXT (ID bản ghi đã xoá)",
        "type: INT (0: subject, 1: document)",
        "dateTimeDeleted: DATETIME"
    ])
    box3 = patches.FancyBboxPatch((0.5, 0.4), 3.8, 1.5, boxstyle="round,pad=0.08", fc='#FFEBEE', ec='#D32F2F', lw=2)
    ax.add_patch(box3)
    ax.text(2.4, 1.6, del_title, ha='center', va='center', fontsize=9.5, fontweight='bold', color='#B71C1C')
    ax.plot([0.5, 4.3], [1.4, 1.4], color='#D32F2F', lw=1.5)
    ax.text(0.7, 0.9, del_fields, ha='left', va='center', fontsize=8, fontfamily='sans-serif', color='#212121')

    # Relationship line 1 - N between Subjects and Documents
    ax.annotate("1", xy=(4.3, 4.2), xytext=(4.45, 4.2), fontsize=12, fontweight='bold', color='#1A237E')
    ax.annotate("N", xy=(5.55, 4.2), xytext=(5.4, 4.2), fontsize=12, fontweight='bold', color='#1A237E')
    ax.annotate("", xy=(5.7, 4.1), xytext=(4.3, 4.1),
                arrowprops=dict(arrowstyle="->", color='#3F51B5', lw=2))

    plt.title("Sơ đồ Thực thể Liên kết (Entity Relationship Diagram - ERD) - SQLite StudyHub", fontsize=13, fontweight='bold', pad=15)
    out_path = os.path.join(IMG_DIR, "erd_database.png")
    plt.tight_layout()
    plt.savefig(out_path, bbox_inches='tight')
    plt.close()
    return out_path

# -------------------------------------------------------------
# 2. HÀM TẠO FILE WORD ĐẸP VÀ CHUYÊN NGHIỆP
# -------------------------------------------------------------

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_styled_table(doc, headers, rows_data):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], '3F51B5')
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)

    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = 'F8F9FA' if r_idx % 2 == 1 else 'FFFFFF'
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(33, 33, 33)

    doc.add_paragraph() # Spacing

def safe_save(doc, default_path):
    try:
        doc.save(default_path)
        return default_path
    except PermissionError:
        base, ext = os.path.splitext(default_path)
        alt_path = f"{base}_CAP_NHAT{ext}"
        try:
            doc.save(alt_path)
            print(f"File {default_path} đang được mở trong Word, đã lưu thành công bản cập nhật tại: {alt_path}")
            return alt_path
        except PermissionError:
            import time
            ts_path = f"{base}_{int(time.time())}{ext}"
            doc.save(ts_path)
            print(f"Lưu thành công tại: {ts_path}")
            return ts_path

# -------------------------------------------------------------
# 3. TẠO FILE WORD: TÀI LIỆU PHÂN TÍCH YÊU CẦU
# -------------------------------------------------------------

def build_requirements_docx(img_dfd0, img_dfd1, img_erd):
    doc = Document()

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("TÀI LIỆU PHÂN TÍCH YÊU CẦU PHẦN MỀM\n(SOFTWARE REQUIREMENTS SPECIFICATION)")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(63, 81, 181)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("DỰ ÁN: STUDYHUB - HỆ THỐNG QUẢN LÝ TÀI LIỆU HỌC TẬP THÔNG MINH\nNhóm phát triển: Nhóm 13 | Môn học: Lập trình Thiết bị Di động (CSE441)")
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Mục 1: Giới thiệu
    h1 = doc.add_heading("1. GIỚI THIỆU TỔNG QUAN HỆ THỐNG", level=1)
    doc.add_paragraph(
        "StudyHub là ứng dụng di động và đa nền tảng (hỗ trợ Android, iOS, Web, Windows) được phát triển nhằm giải quyết "
        "nhu cầu quản lý, tổ chức và tra cứu tài liệu học tập của sinh viên đại học. Ứng dụng tuân thủ nghiêm ngặt chuẩn kiến trúc "
        "Cashew (Local-First Architecture), đảm bảo toàn bộ dữ liệu tài liệu, bài giảng, đề kiểm tra và ghi chú được lưu trữ hoàn toàn "
        "cục bộ trên thiết bị của sinh viên, mang lại tốc độ tức thì và hoạt động 100% khi không có kết nối Internet."
    )

    # Mục 2: Yêu cầu chức năng
    doc.add_heading("2. DANH MỤC YÊU CẦU CHỨC NĂNG (FUNCTIONAL REQUIREMENTS)", level=1)
    
    headers_req = ["Mã YC", "Tên chức năng", "Mô tả chi tiết", "Độ ưu tiên"]
    req_data = [
        ["FR-01", "Quản lý môn học", "Tạo, chỉnh sửa, xóa danh mục môn học với mã môn, tên, màu sắc pastel và biểu tượng trực quan.", "Bắt buộc (High)"],
        ["FR-02", "Quản lý tài liệu đa định dạng", "Thêm, sửa, xóa tài liệu các định dạng: PDF, Word (.docx), PowerPoint (.pptx), Liên kết web (URL).", "Bắt buộc (High)"],
        ["FR-03", "Tự động nhận diện định dạng tệp", "Khi chọn tệp từ máy, hệ thống tự động nhận diện định dạng (Word, PPT, PDF) theo phần mở rộng, loại bỏ sai sót thủ công.", "Bắt buộc (High)"],
        ["FR-04", "Phân loại tài liệu đa mục đích", "Phân loại từng tài liệu theo 4 danh mục chuyên biệt: Bài giảng, Tài liệu tham khảo, Bài tập, Đề kiểm tra.", "Bắt buộc (High)"],
        ["FR-05", "Xem PDF trực tiếp & Tải xuống", "Tích hợp trình xem trước PDF offline tốc độ cao, hỗ trợ thu phóng, cuộn trang mượt mà và nút xuất/tải file về máy.", "Bắt buộc (High)"],
        ["FR-06", "Tìm kiếm & Lọc tài liệu đa tiêu chí", "Tìm kiếm theo từ khóa tên hoặc ghi chú; Lọc nhanh theo Môn học, Loại định dạng, và Phân loại mục đích; Sắp xếp theo ngày/tên.", "Bắt buộc (High)"],
        ["FR-07", "Đánh dấu quan trọng (Ghim)", "Gắn sao một chạm để lưu nhanh các tài liệu quan trọng, đề cương thi cử vào màn hình riêng.", "Bắt buộc (High)"],
        ["FR-08", "Đồng bộ hóa đám mây Google", "Tích hợp đăng nhập Google để đồng bộ CSDL đám mây. Người dùng không đăng nhập vẫn sử dụng bình thường 100% offline.", "Nâng cao (Medium)"],
        ["FR-09", "Xoá an toàn với Tombstone", "Mỗi khi xoá môn học hoặc tài liệu, hệ thống ghi nhận bản ghi huỷ vào bảng DeleteLogs theo chuẩn Cashew.", "Bắt buộc (High)"]
    ]
    create_styled_table(doc, headers_req, req_data)

    # Mục 3: Sơ đồ luồng dữ liệu DFD
    doc.add_heading("3. SƠ ĐỒ LUỒNG DỮ LIỆU (DATA FLOW DIAGRAMS - DFD)", level=1)
    doc.add_paragraph("Dưới đây là sơ đồ luồng dữ liệu minh họa chi tiết dòng lưu chuyển dữ liệu qua các tiến trình và kho lưu trữ của ứng dụng StudyHub:")

    # Hình DFD Mức 0
    doc.add_heading("3.1. Sơ đồ DFD Mức 0 (Context Diagram - Sơ đồ ngữ cảnh)", level=2)
    p_img0 = doc.add_paragraph()
    p_img0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(img_dfd0, width=Inches(6.2))
    p_cap0 = doc.add_paragraph("Hình 1: Sơ đồ luồng dữ liệu DFD Mức 0 tổng thể hệ thống StudyHub")
    p_cap0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap0.runs[0].font.italic = True
    p_cap0.runs[0].font.size = Pt(9.5)

    # Hình DFD Mức 1
    doc.add_heading("3.2. Sơ đồ DFD Mức 1 (Chi tiết 4 tiến trình hệ thống)", level=2)
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(img_dfd1, width=Inches(6.2))
    p_cap1 = doc.add_paragraph("Hình 2: Sơ đồ luồng dữ liệu DFD Mức 1 phân rã tiến trình và kết nối kho dữ liệu")
    p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap1.runs[0].font.italic = True
    p_cap1.runs[0].font.size = Pt(9.5)

    # Mục 4: Sơ đồ CSDL ERD
    doc.add_heading("4. MÔ HÌNH DỮ LIỆU QUAN HỆ (ENTITY RELATIONSHIP DIAGRAM - ERD)", level=1)
    doc.add_paragraph(
        "Cơ sở dữ liệu của StudyHub được quản lý bởi Drift ORM trên nền tảng SQLite Local. "
        "Mô hình bao gồm 3 thực thể cốt lõi: Subjects (Môn học), Documents (Tài liệu) và DeleteLogs (Tombstone đồng bộ hóa):"
    )
    p_erd = doc.add_paragraph()
    p_erd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(img_erd, width=Inches(6.2))
    p_cap_erd = doc.add_paragraph("Hình 3: Sơ đồ thực thể liên kết ERD của ứng dụng StudyHub")
    p_cap_erd.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap_erd.runs[0].font.italic = True
    p_cap_erd.runs[0].font.size = Pt(9.5)

    # Mục 5: Yêu cầu phi chức năng
    doc.add_heading("5. YÊU CẦU PHI CHỨC NĂNG (NON-FUNCTIONAL REQUIREMENTS)", level=1)
    headers_nfr = ["Tiêu chí", "Yêu cầu kỹ thuật đạt được"]
    nfr_data = [
        ["Hiệu năng & Tốc độ", "Thời gian nạp danh sách và lọc < 50ms nhờ Reactive Streams; mở file PDF in-memory tức thì."],
        ["Khả năng Offline-First", "Hoạt động 100% không cần mạng Internet. Không gây phụ thuộc máy chủ."],
        ["Độ tin cậy & Chống sọc vàng", "Kiểm thử nghiêm ngặt không phát sinh lỗi tràn giao diện (Zero RenderFlex Overflow) trên cả mobile lẫn web."],
        ["Giao diện & Trải nghiệm (UI/UX)", "Material 3 chuẩn hoá, hỗ trợ Dark/Light Theme độ tương phản cao, bố cục Grid View hiện đại tự động co giãn."]
    ]
    create_styled_table(doc, headers_nfr, nfr_data)

    out_file = os.path.join(DOCS_DIR, "TAI_LIEU_PHAN_TICH_YEU_CAU.docx")
    return safe_save(doc, out_file)

# -------------------------------------------------------------
# 4. TẠO FILE WORD: BÁO CÁO GIẢI TRÌNH KIẾN TRÚC CASHEW
# -------------------------------------------------------------

def build_architecture_docx(img_arch, img_erd):
    doc = Document()

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("BÁO CÁO GIẢI TRÌNH KIẾN TRÚC CASHEW\nVÀ TRIỂN KHAI HỆ THỐNG STUDYHUB")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(63, 81, 181)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Đối sánh chuẩn mẫu mã nguồn Cashew App & Giải pháp thiết kế Local-First\nNhóm phát triển: Nhóm 13 | Lớp: CSE441")
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Mục 1: Bối cảnh kiến trúc
    doc.add_heading("1. TỔNG QUAN KIẾN TRÚC NỀN TẢNG CASHEW", level=1)
    doc.add_paragraph(
        "Cashew là ứng dụng quản lý tài chính cá nhân mã nguồn mở hàng đầu trên Flutter, nổi tiếng với kiến trúc Local-First, "
        "hiệu năng cực cao và trải nghiệm người dùng mượt mà. Dự án StudyHub kế thừa trọn vẹn triết lý thiết kế và quy chuẩn mã nguồn "
        "của Cashew để xây dựng hệ thống quản lý tài liệu học tập toàn diện cho sinh viên."
    )

    # Mục 2: 5 Tầng kiến trúc
    doc.add_heading("2. SƠ ĐỒ VÀ PHÂN LỚP 5 TẦNG KIẾN TRÚC (5-LAYER ARCHITECTURE)", level=1)
    doc.add_paragraph(
        "Hệ thống StudyHub được phân tách thành 5 phân lớp độc lập, đảm bảo tính đóng gói (encapsulation), "
        "khả năng tái sử dụng mã nguồn và dễ dàng kiểm thử đơn vị (Unit Testable):"
    )

    # Hình kiến trúc 5 tầng
    p_arch = doc.add_paragraph()
    p_arch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(img_arch, width=Inches(6.2))
    p_cap_arch = doc.add_paragraph("Hình 1: Mô hình phân lớp 5 tầng chuẩn Cashew áp dụng trong StudyHub")
    p_cap_arch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap_arch.runs[0].font.italic = True
    p_cap_arch.runs[0].font.size = Pt(9.5)

    headers_layer = ["Tầng kiến trúc", "Thành phần kỹ thuật", "Trách nhiệm chính"]
    layer_data = [
        ["1. Presentation Layer", "DocumentsScreen, SubjectDetailScreen, FavoritesScreen, PdfViewerScreen", "Hiển thị giao diện Material 3, xử lý thao tác chạm, bố cục dạng lưới Grid View co giãn."],
        ["2. State Layer", "ValueNotifier, Reactive Streams, AuthSyncService, Cashew BottomSheets", "Quản lý trạng thái giao diện, kiểm soát đồng bộ Google Account, phản hồi luồng dữ liệu thời gian thực."],
        ["3. Domain Logic", "Formatters, Auto-detect file extension, Pastel Color Algorithm", "Quy tắc nghiệp vụ: tự nhận diện Word/PPT/PDF, mã màu pastel, định dạng dung lượng file bytes."],
        ["4. Data Access (DAO)", "SubjectDao, DocumentDao, FileStorageService, DeleteLogs", "Truy vấn CSDL Drift SQLite, cơ chế Tombstone chống mất mát dữ liệu khi xóa."],
        ["5. Persistence Layer", "AppDatabase (Drift ORM), SQLite Native / Web IndexedDB", "Lưu trữ dữ liệu cục bộ an toàn, cam kết vận hành 100% Local-First không gián đoạn."]
    ]
    create_styled_table(doc, headers_layer, layer_data)

    # Mục 3: Cơ chế Tombstone Soft-Delete
    doc.add_heading("3. CƠ CHẾ TOMBSTONE VÀ TÍNH TOÀN VẸN DỮ LIỆU ĐỒNG BỘ", level=1)
    doc.add_paragraph(
        "Một trong những trụ cột cốt lõi của kiến trúc Cashew là cơ chế Tombstone (nhật ký xoá). Thay vì xoá cứng (hard delete) "
        "mà không để lại dấu vết, hệ thống lưu trữ bản ghi vào bảng DeleteLogs với khóa chính `entryPk`, loại đối tượng và thời điểm xoá. "
        "Nhờ đó, khi người dùng kích hoạt tính năng đồng bộ Google Cloud (qua nút Profile), các thiết bị khác nhận diện chính xác "
        "bản ghi nào cần bị loại bỏ, tránh tình trạng tái sinh dữ liệu đã xoá (data resurrection bug)."
    )

    # Mục 4: Bảng đối sánh trực tiếp với Cashew
    doc.add_heading("4. BẢNG ĐỐI SÁNH TRỰC TIẾP GIỮA CASHEW VÀ STUDYHUB", level=1)
    headers_cmp = ["Tiêu chí kiến trúc", "Triển khai gốc tại Cashew", "Ứng dụng tương đương tại StudyHub"]
    cmp_data = [
        ["Thực thể danh mục", "Ví dụ: Categories, Wallets (Ví, Danh mục thu chi)", "Môn học (Subjects) với mã môn, tên, icon và màu pastel."],
        ["Thực thể chi tiết", "Transactions (Giao dịch thu / chi)", "Documents (Tài liệu học tập) với file PDF, Word, PPT, Link."],
        ["Phân loại mục đích", "Category Type (Thu, Chi, Chuyển khoản)", "Phân loại tài liệu: Bài giảng, Tham khảo, Bài tập, Đề kiểm tra."],
        ["Mã định danh duy nhất", "UUID v4 (String) tự sinh client-side", "TextColumn UUID v4 tự sinh, ngăn ngừa xung đột ID khi sync."],
        ["Phản ứng luồng dữ liệu", "Drift Watch Streams kết hợp StreamBuilder", "StreamBuilder gắn kết chặt chẽ với DocumentDao.watchFilteredDocuments."],
        ["Đồng bộ đám mây", "Google Drive Sync (Offline-First)", "Google Cloud Sync với AuthSyncService, không bắt buộc đăng nhập."]
    ]
    create_styled_table(doc, headers_cmp, cmp_data)

    # Mục 5: Giải pháp chống tràn giao diện (Zero Yellow Stripes)
    doc.add_heading("5. GIẢI PHÁP THIẾT KẾ GIAO DIỆN KHÔNG TRÀN LỖI (ZERO YELLOW STRIPES)", level=1)
    doc.add_paragraph(
        "Nhằm loại bỏ hoàn toàn các lỗi sọc vàng đen (RenderFlex overflow) phổ biến trên Flutter, StudyHub đã chuẩn hóa:\n"
        "1. Thẻ tài liệu (DocumentCard): Cố định 4 slot hành động bằng IconButton thẳng hàng (Xem trước, Tải xuống, Sửa, Xoá) với chiều rộng vừa vặn 160px; khi không xem trước được thì giữ khoảng trống để 3 icon còn lại không co giãn.\n"
        "2. Trường Môn học: Bổ sung thuộc tính `isExpanded: true` và `TextOverflow.ellipsis`, loại bỏ tràn chiều ngang khi tên môn học dài.\n"
        "3. Trường Ghi chú: Placeholder căn giữa dọc khi rỗng và tự động co giãn lên đến 4 dòng khi người dùng nhập nội dung dài.\n"
        "4. Bố cục Grid View: Co giãn linh hoạt `maxCrossAxisExtent: 460`, bảo đảm giao diện chuẩn mực trên mọi độ phân giải thiết bị."
    )

    out_file = os.path.join(DOCS_DIR, "BAO_CAO_GIAI_TRINH_KIEN_TRUC_CASHEW.docx")
    return safe_save(doc, out_file)

def main():
    print("Bắt đầu sinh ảnh sơ đồ...")
    img0 = draw_dfd_level_0()
    img1 = draw_dfd_level_1()
    img_arch = draw_cashew_architecture()
    img_erd = draw_erd()
    print("Đã tạo ảnh thành công!")

    print("Bắt đầu tạo tài liệu Word TAI_LIEU_PHAN_TICH_YEU_CAU.docx...")
    doc1 = build_requirements_docx(img0, img1, img_erd)
    print(f"-> Tạo thành công: {doc1}")

    print("Bắt đầu tạo tài liệu Word BAO_CAO_GIAI_TRINH_KIEN_TRUC_CASHEW.docx...")
    doc2 = build_architecture_docx(img_arch, img_erd)
    print(f"-> Tạo thành công: {doc2}")

if __name__ == '__main__':
    main()
